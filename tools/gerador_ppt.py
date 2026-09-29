import os
try:
    from pptx import Presentation
except ImportError:
    print("Atenção: A biblioteca python-pptx não foi encontrada.")

def criar_apresentacao(tema, texto_slides):
    try:
        prs = Presentation()
        
        # Slide de Título
        slide_titulo = prs.slides.add_slide(prs.slide_layouts[0])
        if slide_titulo.shapes.title:
            slide_titulo.shapes.title.text = tema.title()
        if len(slide_titulo.placeholders) > 1:
            slide_titulo.placeholders[1].text = "Gerado pelo Jarvis"
        
        # Limpa o texto para evitar falhas de formatação do LLM
        texto_limpo = texto_slides.replace("###", "").replace("**", "")
        blocos = texto_limpo.split('\n\n')
        
        for bloco in blocos:
            if bloco.strip():
                slide = prs.slides.add_slide(prs.slide_layouts[1])
                linhas = bloco.strip().split('\n')
                
                # A primeira linha vira o título do slide
                if slide.shapes.title:
                    slide.shapes.title.text = linhas[0].strip()
                
                # O restante vira o corpo de texto
                if len(slide.placeholders) > 1:
                    corpo = slide.placeholders[1].text_frame
                    corpo.text = '\n'.join(linhas[1:]).replace("*", "-") if len(linhas) > 1 else ""
                
        nome_arquivo = f"{tema.replace(' ', '_')}.pptx"
        caminho_arquivo = os.path.join(os.getcwd(), nome_arquivo)
        
        # Salva o arquivo fisicamente
        prs.save(caminho_arquivo)
        return caminho_arquivo
        
    except Exception as e:
        print(f"\n[ERRO FATAL NO PPT]: {e}\n") # Agora o erro vai gritar no seu terminal!
        return f"Erro: {e}"